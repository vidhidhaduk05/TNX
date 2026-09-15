"""Render the combined Table S1 to a house-style Word .docx.

Matches the delivered cSDH docs' house style:
  - Page: Tabloid (11 x 17 in), portrait, narrow margins
  - Font: Times New Roman 11 pt
  - Bold "Supplementary Table S1:" title spanning the table width
  - Header row bold; section rows (Inclusion/Outcomes/Covariates) bold+italic, shaded
  - Thin cell borders throughout; a vertical rule separating the label column
  - 3 columns: Characteristic | Codes | Applies to
"""
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

sys.path.insert(0, "/workspace")
from build_combined_s1 import build_merged, studies_label, STUDY_ORDER  # noqa: E402

FONT = "Times New Roman"
FONT_PT = 11
TITLE_PT = 11
GRID = "BBBBBB"
SECTION_SHADE = "D9D9D9"


def set_cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        spec = kw.get(edge)
        if spec is None:
            continue
        el = borders.find(qn("w:" + edge))
        if el is None:
            el = OxmlElement("w:" + edge)
            borders.append(el)
        for k, v in spec.items():
            el.set(qn("w:" + k), str(v))


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def style_run(run, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(FONT_PT)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), FONT)


def write_cell(cell, text, bold=False, italic=False):
    p = cell.paragraphs[0]
    # remove any pre-existing runs so the cell has exactly one clean styled run
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(text)
    style_run(run, bold=bold, italic=italic)
    thin = {"val": "single", "sz": 4, "color": GRID, "space": 0}
    set_cell_border(cell, top=thin, left=thin, bottom=thin, right=thin)


def build_docx(out_path):
    rows = build_merged()

    doc = Document()
    sec = doc.sections[0]
    # Tabloid portrait
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width = Inches(11)
    sec.page_height = Inches(17)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Inches(0.6))

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(FONT_PT)

    # Title paragraph
    title = ("Supplementary Table S1: Diagnosis, procedure, medication, and laboratory "
             "codes used across all five chronic subdural hematoma (cSDH) analyses")
    p = doc.add_paragraph()
    r = p.add_run("Supplementary Table S1: ")
    style_run(r, bold=True)
    r2 = p.add_run(title.split(": ", 1)[1])
    style_run(r2, bold=False)

    ncols = 3
    table = doc.add_table(rows=1, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(3.6), Inches(5.2), Inches(1.4)]

    hdr = ["Characteristic", "Codes (ICD-10-CM / ICD-10-PCS / RxNorm / TriNetX)", "Applies to"]
    for c, h in enumerate(hdr):
        write_cell(table.rows[0].cells[c], h, bold=True)

    cur_sec = None
    for row in rows:
        if row["section"] != cur_sec:
            cur_sec = row["section"]
            rc = table.add_row().cells
            write_cell(rc[0], cur_sec, bold=True, italic=True)
            write_cell(rc[1], "", bold=True, italic=True)
            write_cell(rc[2], "", bold=True, italic=True)
            for c in range(ncols):
                shade_cell(rc[c], SECTION_SHADE)
        rc = table.add_row().cells
        write_cell(rc[0], row["label"])
        write_cell(rc[1], row["code"])
        write_cell(rc[2], studies_label(row["studies"]))

    # apply column widths to every cell
    for r_ in table.rows:
        for c, w in enumerate(widths):
            r_.cells[c].width = w

    # thicker vertical rule right of label column
    thickish = {"val": "single", "sz": 8, "color": "000000", "space": 0}
    for r_ in table.rows:
        set_cell_border(r_.cells[0], right=thickish)
        set_cell_border(r_.cells[1], left=thickish)

    doc.save(out_path)
    return len(table.rows)


if __name__ == "__main__":
    n = build_docx("/workspace/combined_S1.docx")
    print(f"wrote /workspace/combined_S1.docx with {n} table rows (incl. header + section rows)")
